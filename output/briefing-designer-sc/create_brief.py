from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from pathlib import Path
out=Path('/Users/Shared/Projects/RunnerHub/Business/output/briefing-designer-sc')
d=Document(); sec=d.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.8); sec.left_margin=Cm(2); sec.right_margin=Cm(2)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 s=d.styles[name]; s.font.name='Arial'; s.font.color.rgb=RGBColor(0,0,0)
s=d.styles['Normal']; s.font.size=Pt(10.5); s.paragraph_format.space_after=Pt(6); s.paragraph_format.line_spacing=1.08
for name,size in [('Title',23),('Heading 1',14),('Heading 2',11)]:
 s=d.styles[name]; s.font.size=Pt(size); s.font.bold=True
 s.paragraph_format.space_before=Pt(12 if name!='Title' else 0); s.paragraph_format.space_after=Pt(6)
d.core_properties.title='Briefing de peças para Display RoadRunners em Santa Catarina'
d.core_properties.author='RoadRunners'
def p(t,style=None): return d.add_paragraph(t,style)
def h(t): return d.add_heading(t,1)
def link(label,url):
 a=d.add_paragraph(); hl=OxmlElement('w:hyperlink'); hl.set(qn('r:id'),a.part.relate_to(url,RT.HYPERLINK,is_external=True)); r=OxmlElement('w:r'); pr=OxmlElement('w:rPr'); c=OxmlElement('w:color'); c.set(qn('w:val'),'0563C1'); pr.append(c); r.append(pr); t=OxmlElement('w:t'); t.text=label; r.append(t); hl.append(r); a._p.append(hl)
p('Briefing de peças para Display', 'Title')
p('RoadRunners em Santa Catarina', 'Subtitle')
p('Para a equipe de design • 1 de outubro de 2026')
p('Criar 2 conceitos visuais, cada um em 3 formatos, e exportar 2 versões do logotipo oficial. A entrega terá 8 arquivos finais para anúncios responsivos do Google Display, com foco no benefício de encontrar a próxima corrida em Santa Catarina.')
h('Objetivo e público')
p('Atrair corredores de Santa Catarina para consultar o calendário e escolher provas por cidade, data e distância. A comunicação deve atender tanto quem está começando quanto quem já participa de corridas.')
p('Mensagem principal: Encontre sua próxima corrida em SC.')
link('Destino da campanha — calendário de Santa Catarina','https://roadrunners.run/estado/sc/')
h('Quantidade e dimensões')
t=d.add_table(rows=1, cols=4); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
widths=[5.4,4.3,3.2,4.1]
for c,w in zip(t.columns,widths): c.width=Cm(w)
rows=[['Arquivo','Dimensões em pixels','Proporção','Quantidade'],['Imagem horizontal','1200 × 628','1,91:1','2 • uma por conceito'],['Imagem quadrada','1200 × 1200','1:1','2 • uma por conceito'],['Imagem vertical','900 × 1600','9:16','2 • uma por conceito'],['Logotipo quadrado','1200 × 1200','1:1','1'],['Logotipo horizontal','1200 × 300','4:1','1']]
for i,vals in enumerate(rows):
 cells=t.rows[0].cells if i==0 else t.add_row().cells
 for j,(c,txt) in enumerate(zip(cells,vals)):
  c.width=Cm(widths[j]); c.text=txt; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  tcpr=c._tc.get_or_add_tcPr(); margins=OxmlElement('w:tcMar')
  for side in ['top','left','bottom','right']:
   e=OxmlElement('w:'+side); e.set(qn('w:w'),'100'); e.set(qn('w:type'),'dxa'); margins.append(e)
  tcpr.append(margins); shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E9E9E9' if i==0 else 'FFFFFF'); tcpr.append(shade)
  borders=OxmlElement('w:tcBorders')
  for side in ['top','left','bottom','right']:
   e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
  tcpr.append(borders)
  for para in c.paragraphs:
   para.paragraph_format.space_after=Pt(0)
   if j==2: para.alignment=WD_ALIGN_PARAGRAPH.CENTER
   for r in para.runs: r.font.size=Pt(9); r.bold=i==0
 rep=OxmlElement('w:cantSplit'); t.rows[i]._tr.get_or_add_trPr().append(rep)
p('Total: 6 imagens + 2 logotipos. Priorizar as versões horizontal e quadrada; as verticais complementam o pacote. Os logotipos podem ser exportados do arquivo oficial existente.')
h('Dois conceitos visuais')
p('Conceito 1 — Sua próxima largada. Fotografia de corrida de rua com participação e movimento. Transmitir a vontade de escolher uma prova e participar, sem vincular a peça a uma única organização ou evento.')
p('Conceito 2 — Do treino à próxima prova. Pessoa ou pequeno grupo correndo em parque ou ambiente urbano, com expressão natural. Aproximar o treino cotidiano da descoberta de um próximo desafio.')
d.add_page_break()
h('Direção de arte')
p('Usar a identidade oficial do RoadRunners, com preto, branco e amarelo como referências. Valorizar corredores reais e diversidade de perfis, com fotografia natural e boa leitura em formatos pequenos.')
p('Preferir fotos próprias ou licenciadas de Santa Catarina. Se houver um local reconhecível, confirmar que fica em SC. Uma paisagem neutra também funciona. Não usar pontos turísticos do Rio de Janeiro ou de outros estados como ambientação desta campanha.')
p('Entregar as imagens sem textos, logotipos ou botões sobrepostos. No anúncio responsivo, o Google combina a fotografia com o logotipo e os textos enviados separadamente. Não montar colagens nem simular uma interface clicável.')
p('Adaptar o enquadramento de cada formato, sem esticar a foto. Manter rostos e elementos essenciais longe das bordas para tolerar recortes. Evitar que marcas de roupas, patrocinadores ou a LIVE! dominem a composição.')
h('Textos de referência do anúncio')
p('Os textos abaixo contextualizam a criação. Devem permanecer separados das imagens.')
p('Títulos curtos', 'Heading 2')
for txt in ['Sua próxima corrida em SC','Encontre provas na sua região','Calendário de corridas SC']: p('• '+txt)
p('Título longo', 'Heading 2')
p('Encontre sua próxima corrida em Santa Catarina. Consulte cidades, datas e distâncias.')
p('Descrições', 'Heading 2')
p('Explore o calendário de corridas em SC e veja os detalhes das provas no Road Runners.')
p('Escolha seu próximo desafio. Encontre eventos por cidade e distância.')
h('Entrega dos arquivos')
p('Exportar as 6 imagens em JPG ou PNG, com até 5.120 KB por arquivo. Para os 2 logotipos, preferir PNG com fundo transparente e garantir leitura nas duas proporções. Enviar também os arquivos editáveis e as fotos originais usadas, com a origem ou licença de uso.')
p('Nomear por conceito e dimensão, por exemplo: rr_sc_largada_1200x628.jpg e rr_sc_treino_900x1600.jpg. Conferir nitidez, proporção e recortes antes de entregar.')
h('Referências técnicas')
link('Google Ads — formatos de imagens e logotipos','https://support.google.com/google-ads/answer/7005917')
link('Google Ads — boas práticas para anúncios responsivos de Display','https://support.google.com/google-ads/answer/9823397?hl=en')
d.save(out/'Briefing_RoadRunners_Display_SC.docx')
print(out/'Briefing_RoadRunners_Display_SC.docx')
